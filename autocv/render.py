"""Render a TailoredResume to .docx in one of the themes (themes.py).

Direct formatting (no Word styles): each paragraph type is encoded as pPr/rPr XML. The
Classic theme reproduces the original resume's design exactly. The template
(data/templates/base.docx) carries the Word theme, settings and page setup with an empty
body and no personal data; page size and margins come from the theme and paper setting.
"""

from __future__ import annotations

import datetime as dt
import re
from contextvars import ContextVar
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml

from . import themes
from .schema import MasterProfile, TailoredResume
from .themes import Theme

TEMPLATE = Path(__file__).resolve().parent / "data" / "templates" / "base.docx"

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
R_NS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

SECTION_TITLES = themes.CLASSIC.titles  # the Classic headings (what `autocv ingest` reads)

_SPACING = '<w:spacing w:before="{before}" w:after="{after}" w:line="240" w:lineRule="auto"/>'


def _rpr(color: str, size: int, *, bold=False, italic=False, font="Calibri") -> str:
    return (
        f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}"/>'
        + ("<w:b/>" if bold else "")
        + ("<w:i/>" if italic else "")
        + f'<w:color w:val="{color}"/><w:sz w:val="{size}"/></w:rPr>'
    )


_XML_INVALID = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")  # e.g. form feeds from PDFs


def clean_text(text: str) -> str:
    return _XML_INVALID.sub(" ", text or "")


class _Builder:
    """Paragraph types of the resume design, encoded as direct formatting (no Word styles)."""

    def __init__(self, template: Path, theme: Theme, paper: str | None):
        self.t = theme
        self.doc = Document(str(template))
        self.body = self.doc.element.body
        self.sect = self.body[-1]  # sectPr stays last
        width, height = theme.page(paper)
        top, right, bottom, left = theme.margins
        _set(self.sect.find(f"{W}pgSz"), w=width, h=height)
        _set(self.sect.find(f"{W}pgMar"), top=top, right=right, bottom=bottom, left=left)
        self._rule = f'<w:pBdr><w:bottom w:val="single" w:sz="4" w:space="2" w:color="{theme.rule}"/></w:pBdr>'
        self._hang = f'<w:ind w:left="{theme.hang}" w:hanging="{theme.hang}"/>'
        self._tab = f'<w:tabs><w:tab w:val="right" w:pos="{width - left - right}"/></w:tabs>'

    def _run(self, text: str, color: str, size: int, *, font: str | None = None, **kw) -> str:
        rpr = _rpr(color, size, font=font or self.t.body_font, **kw)
        return f'<w:r>{rpr}<w:t xml:space="preserve">{escape(clean_text(text))}</w:t></w:r>'

    def _para(self, runs: str, *, before=0, after=None, rule=False, hang=False, tab=False, keep_next=False) -> str:
        # keep_next: headings stay on the same page as the text under them (never stranded at a page bottom)
        after = self.t.after if after is None else after
        spacing = _SPACING.format(before=before, after=after).replace(' w:before="0"', "")
        ppr = ("<w:keepNext/>" if keep_next else "") + "<w:keepLines/>" + (self._tab if tab else "") \
            + (self._rule if rule else "") + spacing
        ppr += self._hang if hang else ""
        return f"<w:p {W_NS} {R_NS}><w:pPr>{ppr}</w:pPr>{runs}</w:p>"

    def add(self, xml: str) -> None:
        self.sect.addprevious(parse_xml(xml))

    def hyperlink(self, text: str, url: str) -> str:
        rid = self.doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
        inner = self._run(text, self.t.accent, self.t.contact_size)
        return f'<w:hyperlink r:id="{rid}">{inner}</w:hyperlink>'

    # paragraph types ------------------------------------------------------------
    def name(self, text):
        t = self.t
        self.add(self._para(self._run(text, t.ink, t.name_size, bold=True, font=t.name_font), after=t.name_after))

    def headline(self, text):
        self.add(self._para(self._run(text, self.t.accent, self.t.headline_size, bold=True), after=self.t.name_after))

    def contact(self, parts: list[tuple[str, str | None]]):
        t, runs = self.t, []
        for i, (text, url) in enumerate(parts):
            if i:
                runs.append(self._run(t.separator, t.muted, t.contact_size))
            runs.append(self.hyperlink(text, url) if url else self._run(text, t.muted, t.contact_size))
        self.add(self._para("".join(runs), after=t.contact_after))

    def section(self, title):
        t = self.t
        self.add(self._para(self._run(title, t.accent, t.section_size, bold=True), before=t.section_before,
                            after=t.section_after, rule=True, keep_next=True))

    def summary(self, text):
        # The original summary shares the heading's bottom rule (Word draws one rule
        # under the grouped pair), so the paragraph properties are the same.
        t = self.t
        self.add(self._para(self._run(text, t.body, t.body_size), before=t.section_before, after=t.section_after,
                            rule=True))

    def bullet(self, text):
        self.add(self._para(self._run(self.t.bullet + text, self.t.body, self.t.body_size), hang=True))

    def lead_bullet(self, label, text):
        t = self.t
        runs = (self._run(t.bullet, t.body, t.body_size) + self._run(label, t.ink, t.body_size, bold=True)
                + self._run(text, t.body, t.body_size))
        self.add(self._para(runs, hang=True))

    def company(self, left, dates):
        t = self.t
        runs = self._run(left, t.ink, t.company_size, bold=True) + "<w:r><w:tab/></w:r>" + \
            self._run(dates, t.date, t.date_size)
        self.add(self._para(runs, before=t.company_before, after=t.company_after, tab=True, keep_next=True))

    def title(self, text):
        self.add(self._para(self._run(text, self.t.accent, self.t.body_size, bold=True), keep_next=True))

    def scope(self, text, italic=True):
        t = self.t
        self.add(self._para(self._run(text, t.scope if italic else t.body, t.body_size, italic=italic),
                            after=t.scope_after, keep_next=True))


def _set(el, **attrs) -> None:
    for k, v in attrs.items():
        el.set(f"{W}{k}", str(v))


# The design for renders in this request/task: (theme, paper). Set from Settings by the API/CLI.
_DESIGN: ContextVar[tuple[Theme | None, str | None]] = ContextVar("autocv_design", default=(None, None))


def use_design(theme: str | None, paper: str | None):
    return _DESIGN.set((themes.get(theme), paper if paper in themes.PAPER else None))


def active_design() -> tuple[Theme, str]:
    """The theme and paper renders use right now (Classic on its own paper by default)."""
    theme, paper = _DESIGN.get()
    theme = theme or themes.CLASSIC
    return theme, paper or theme.paper


def _lead_text(text: str) -> str:
    return " " + text if text and not text.startswith((" ", ",", ":")) else text


def render(profile: MasterProfile, tailored: TailoredResume, out: Path,
           template: Path = TEMPLATE, theme: Theme | str | None = None, paper: str | None = None) -> Path:
    """Render with a theme (default: the active one, see `use_design`) on Letter or A4."""
    active_theme, active_paper = _DESIGN.get()
    if theme is None:
        theme = active_theme or themes.CLASSIC
    elif isinstance(theme, str):
        theme = themes.get(theme)
    paper = paper or active_paper
    titles = theme.titles
    b = _Builder(template, theme, paper)
    c = profile.contact

    b.name(c.name)
    b.headline(profile.headline(tailored.headline).text)
    parts: list[tuple[str, str | None]] = [(c.location, None)]
    if c.phone:
        parts.append((c.phone, None))
    if c.email:
        parts.append((c.email, f"mailto:{c.email}"))
    parts += [(link.text, link.url) for link in c.links]
    b.contact(parts)

    if tailored.summary:
        b.section(titles["summary"])
        b.summary(tailored.summary.text)

    if tailored.highlights:
        b.section(titles["highlights"])
        for h in tailored.highlights:
            b.bullet(h.text)

    if any(c.items for c in tailored.competencies):
        b.section(titles["competencies"])
        for comp in tailored.competencies:
            if comp.items:  # an emptied group is simply left out
                b.lead_bullet(f"{comp.label}: ", theme.separator.join(comp.items))

    if tailored.experience:
        b.section(titles["experience"])
        for tr in tailored.experience:
            role = profile.role(tr.role)
            b.company(f"{role.employer}, {role.location}", role.dates)
            b.title(role.title)
            if tr.scope:
                b.scope(tr.scope.text, italic=role.scope.italic if role.scope else True)
            for bl in tr.bullets:
                b.bullet(bl.text)
            for sr in tr.sub_roles:
                item = profile.lead_item(sr.id)
                b.lead_bullet(item.label, _lead_text(sr.text.text if sr.text else item.text))

    if tailored.projects:
        b.section(titles["projects"])
        for tp in tailored.projects:
            item = profile.lead_item(tp.id)
            b.lead_bullet(item.label, _lead_text(tp.text.text if tp.text else item.text))

    for key in ("education", "extras"):
        ids = getattr(tailored, key)
        if ids:
            b.section(titles[key])
            for item_id in ids:
                item = profile.lead_item(item_id)
                b.lead_bullet(item.label, _lead_text(item.text))

    props = b.doc.core_properties
    props.author = props.last_modified_by = c.name
    props.title = f"{c.name} — Resume"
    props.created = props.modified = dt.datetime.now(dt.UTC).replace(tzinfo=None)
    out.parent.mkdir(parents=True, exist_ok=True)
    b.doc.save(str(out))
    return out


def docx_text(path: Path) -> list[str]:
    """Paragraph texts of a .docx (tabs kept), for fidelity comparisons."""
    doc = Document(str(path))
    lines = []
    for p in doc.paragraphs:
        text = "".join(
            (node.text or "") if node.tag.endswith("}t") else "\t"
            for node in p._p.iter()
            if node.tag.endswith("}t") or node.tag.endswith("}tab") and node.getparent().tag.endswith("}r")
        )
        lines.append(text)
    return [line for line in lines if line.strip()]
